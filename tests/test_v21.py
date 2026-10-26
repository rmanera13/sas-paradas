from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from app.db import get_db


def _parada(app, eq, motivo_cod, ini_h_atras, dur_min):
    """Insere uma parada fechada diretamente no banco (início X horas atrás)."""
    ini = datetime.now(timezone.utc).replace(microsecond=0) - timedelta(hours=ini_h_atras)
    fim = ini + timedelta(minutes=dur_min)
    with app.app_context():
        db = get_db()
        uid = db.one("INSERT INTO usuario (nome) VALUES ('Teste') ON CONFLICT DO NOTHING RETURNING id") \
            or db.one("SELECT id FROM usuario WHERE nome = 'Teste'")
        mid = db.one("SELECT id FROM motivo WHERE codigo = ?", (motivo_cod,))["id"]
        db.execute("INSERT INTO parada (equipamento_id, usuario_id, motivo_id, inicio, fim, duracao_min, status) "
                   "VALUES (?, ?, ?, ?, ?, ?, 'FECHADA')",
                   (eq, uid["id"], mid, ini.isoformat(), fim.isoformat(), dur_min))
        db.commit()


def test_mtbf_sem_falhas(client):
    r = client.get("/api/v1/analytics/mtbf?dias=7").get_json()
    assert r["dias"] == 7 and r["geral"]["falhas"] == 0 and r["geral"]["mtbf_h"] is None
    assert len(r["por_equipamento"]) == 5


def test_mtbf_calculo(app, client):
    # Máquina 1: 2 falhas de 60 min + 1 preventiva (programada) de 120 min em 7 dias (168 h)
    _parada(app, 1, "MAN-01", 50, 60)
    _parada(app, 1, "OPE-02", 20, 60)
    _parada(app, 1, "MAN-03", 10, 120)
    r = client.get("/api/v1/analytics/mtbf?dias=7").get_json()
    m1 = r["por_equipamento"][0]
    assert m1["falhas"] == 2                       # programada não conta como falha
    assert m1["horas_paradas"] == 4.0              # mas desconta do tempo em operação
    assert m1["horas_operacao"] == 164.0
    assert m1["mtbf_h"] == 82.0                    # 164 h / 2 falhas
    assert m1["mttr_min"] == 60.0
    assert r["geral"]["falhas"] == 2
    assert r["geral"]["mtbf_h"] == round((168 * 5 - 4) / 2, 2)


def test_mtbf_fora_do_periodo(app, client):
    _parada(app, 2, "MAN-01", 24 * 10, 30)         # 10 dias atrás
    assert client.get("/api/v1/analytics/mtbf?dias=7").get_json()["geral"]["falhas"] == 0
    assert client.get("/api/v1/analytics/mtbf?dias=30").get_json()["geral"]["falhas"] == 1


def test_export_excel(app, client):
    _parada(app, 1, "MAN-01", 5, 45)
    r = client.get("/api/v1/export.xlsx?dias=30")
    assert r.status_code == 200
    assert r.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "paradas_" in r.headers["Content-Disposition"] and ".xlsx" in r.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(r.data))
    assert wb.sheetnames == ["Paradas", "MTBF", "Pareto", "Resumo"]
    ws = wb["Paradas"]
    assert ws["B2"].value == "Máquina 01" and ws["F2"].value == 45
    assert isinstance(ws["D2"].value, datetime)    # data real no Excel, não texto
    assert wb["MTBF"]["A7"].value == "TOTAL" and wb["MTBF"]["D7"].value == 1


def test_export_excel_vazio(client):
    wb = load_workbook(BytesIO(client.get("/api/v1/export.xlsx").data))
    assert wb["Paradas"].max_row == 1


def test_dashboard_tem_mtbf_e_excel(client):
    html = client.get("/dashboard").get_data(as_text=True)
    assert "MTBF por máquina" in html and "export.xlsx" in html and ".json" not in html
