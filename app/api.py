from datetime import datetime, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from flask import Blueprint, jsonify, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import get_db, INTEGRITY_ERRORS
from .versao import __version__

bp = Blueprint("api", __name__, url_prefix="/api/v1")

FUSO = ZoneInfo("America/Sao_Paulo")

SELECT_PARADA = """
SELECT p.id, p.equipamento_id, e.nome AS equipamento, u.nome AS operador,
       p.motivo_id, m.codigo AS motivo_codigo, m.descricao AS motivo,
       m.categoria, m.programado, p.inicio, p.fim, p.duracao_min,
       p.observacao, p.status
FROM parada p
JOIN equipamento e ON e.id = p.equipamento_id
JOIN usuario u     ON u.id = p.usuario_id
LEFT JOIN motivo m ON m.id = p.motivo_id
"""


def agora():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def erro(msg, code=400):
    return jsonify({"erro": msg}), code


def buscar(pid):
    return get_db().one(SELECT_PARADA + " WHERE p.id = ?", (pid,))


def operador_id(nome):
    """Operador e digitado: cria o cadastro na primeira vez e reaproveita depois."""
    db = get_db()
    achou = db.one("SELECT id FROM usuario WHERE lower(nome) = lower(?)", (nome,))
    if achou:
        return achou["id"]
    return db.one("INSERT INTO usuario (nome) VALUES (?) RETURNING id", (nome,))["id"]


@bp.get("/health")
def health():
    db = get_db()
    db.one("SELECT 1 AS ok")
    return jsonify({"status": "ok", "versao": __version__,
                    "banco": "postgresql" if db.pg else "sqlite"})


@bp.get("/equipamentos")
def equipamentos():
    rows = get_db().all("""
        SELECT e.id, e.nome,
               (SELECT p.id FROM parada p WHERE p.equipamento_id = e.id AND p.fim IS NULL) AS parada_aberta_id
        FROM equipamento e WHERE e.ativo = 1 ORDER BY e.nome""")
    return jsonify(rows)


@bp.get("/motivos")
def motivos():
    return jsonify(get_db().all("SELECT * FROM motivo ORDER BY codigo"))


@bp.get("/operadores")
def operadores():
    rows = get_db().all("SELECT nome FROM usuario ORDER BY nome")
    return jsonify([r["nome"] for r in rows])


def _filtrar_paradas(args):
    sql, valores = SELECT_PARADA + " WHERE 1=1", []
    if args.get("status"):
        sql += " AND p.status = ?"; valores.append(args["status"])
    if args.get("equipamento"):
        sql += " AND p.equipamento_id = ?"; valores.append(args["equipamento"])
    if args.get("desde"):
        sql += " AND p.inicio >= ?"; valores.append(args["desde"])
    if args.get("ate"):
        sql += " AND p.inicio <= ?"; valores.append(args["ate"])
    sql += " ORDER BY p.inicio DESC, p.id DESC"
    return get_db().all(sql, valores)


@bp.get("/paradas")
def listar():
    return jsonify(_filtrar_paradas(request.args))


@bp.post("/paradas")
def abrir():
    dados = request.get_json(silent=True) or {}
    eq = dados.get("equipamento_id")
    nome = (dados.get("operador") or "").strip()
    if not eq:
        return erro("Informe o equipamento.")
    if len(nome) < 2:
        return erro("Digite o nome do operador (mínimo 2 caracteres).")
    if len(nome) > 80:
        return erro("Nome do operador muito longo (máximo 80).")
    db = get_db()
    if not db.one("SELECT 1 AS ok FROM equipamento WHERE id = ? AND ativo = 1", (eq,)):
        return erro("Equipamento não encontrado.", 404)
    try:
        novo = db.one(
            "INSERT INTO parada (equipamento_id, usuario_id, inicio, status) "
            "VALUES (?, ?, ?, 'ABERTA') RETURNING id",
            (eq, operador_id(nome), agora()))
        db.commit()
    except INTEGRITY_ERRORS:
        db.rollback()
        return erro("Já existe uma parada aberta neste equipamento.", 409)
    return jsonify(buscar(novo["id"])), 201


@bp.patch("/paradas/<int:pid>")
def classificar(pid):
    p = buscar(pid)
    if not p:
        return erro("Parada não encontrada.", 404)
    if p["status"] == "FECHADA":
        return erro("Parada já fechada.", 409)
    dados = request.get_json(silent=True) or {}
    mid = dados.get("motivo_id")
    db = get_db()
    if not mid or not db.one("SELECT 1 AS ok FROM motivo WHERE id = ?", (mid,)):
        return erro("Informe um motivo válido.")
    novo = "FECHADA" if p["status"] == "AGUARDANDO_MOTIVO" else "ABERTA"
    db.execute("UPDATE parada SET motivo_id = ?, observacao = ?, status = ? WHERE id = ?",
               (mid, (dados.get("observacao") or "").strip()[:500] or None, novo, pid))
    db.commit()
    return jsonify(buscar(pid))


@bp.post("/paradas/<int:pid>/retomar")
def retomar(pid):
    p = buscar(pid)
    if not p:
        return erro("Parada não encontrada.", 404)
    if p["fim"]:
        return erro("Parada já retomada.", 409)
    fim = agora()
    dur = round((datetime.fromisoformat(fim) - datetime.fromisoformat(p["inicio"])).total_seconds() / 60, 2)
    status = "FECHADA" if p["motivo_id"] else "AGUARDANDO_MOTIVO"
    db = get_db()
    db.execute("UPDATE parada SET fim = ?, duracao_min = ?, status = ? WHERE id = ?",
               (fim, dur, status, pid))
    db.commit()
    return jsonify(buscar(pid))


@bp.get("/analytics/pareto")
def pareto():
    return jsonify(_pareto())


def _pareto():
    rows = get_db().all("""
        SELECT m.codigo, m.descricao, m.programado,
               ROUND(CAST(SUM(p.duracao_min) AS NUMERIC), 2) AS minutos,
               COUNT(*) AS ocorrencias
        FROM parada p JOIN motivo m ON m.id = p.motivo_id
        WHERE p.status = 'FECHADA'
        GROUP BY m.id, m.codigo, m.descricao, m.programado ORDER BY minutos DESC""")
    total = sum(r["minutos"] or 0 for r in rows)
    acum, out = 0, []
    for i, r in enumerate(rows, 1):
        acum += r["minutos"] or 0
        d = dict(r)
        d["pct_acumulado"] = round(acum / total * 100, 1) if total else round(i / len(rows) * 100, 1)
        out.append(d)
    return out


@bp.get("/analytics/resumo")
def resumo():
    return jsonify(_resumo())


def _resumo():
    db = get_db()
    d = db.one("""
        SELECT COUNT(*) AS total_paradas,
               COALESCE(ROUND(CAST(SUM(p.duracao_min) AS NUMERIC), 2), 0) AS minutos_total,
               COALESCE(ROUND(CAST(AVG(p.duracao_min) AS NUMERIC), 2), 0) AS mttr_min,
               COALESCE(ROUND(CAST(SUM(CASE WHEN m.programado = 1 THEN p.duracao_min END) AS NUMERIC), 2), 0) AS min_programado,
               COALESCE(ROUND(CAST(SUM(CASE WHEN m.programado = 0 THEN p.duracao_min END) AS NUMERIC), 2), 0) AS min_nao_programado
        FROM parada p JOIN motivo m ON m.id = p.motivo_id
        WHERE p.status = 'FECHADA'""")
    d["paradas_abertas"] = db.one(
        "SELECT COUNT(*) AS n FROM parada WHERE status <> 'FECHADA'")["n"]
    return d


# ---------------------------------------------------------------- MTBF (v2.1)

def _periodo(args):
    """Período de análise: ?dias=N (padrão 30) terminando agora."""
    try:
        dias = int(args.get("dias", 30))
    except ValueError:
        dias = 30
    dias = max(1, min(dias, 365))
    fim = datetime.now(timezone.utc).replace(microsecond=0)
    return fim - timedelta(days=dias), fim, dias


def _mtbf(args):
    """MTBF = tempo em operação / nº de falhas (paradas NÃO programadas) no período.

    - Tempo em operação = horas do período - horas paradas (qualquer parada, recortada no período;
      paradas ainda abertas contam até agora).
    - Falha = parada com motivo NÃO programado iniciada no período.
      Paradas ainda sem motivo não entram como falha até serem classificadas.
    """
    ini, fim, dias = _periodo(args)
    horas_periodo = (fim - ini).total_seconds() / 3600
    db = get_db()
    eqs = db.all("SELECT id, nome FROM equipamento WHERE ativo = 1 ORDER BY nome")
    paradas = db.all("""
        SELECT p.equipamento_id, p.inicio, p.fim, p.duracao_min, m.programado
        FROM parada p LEFT JOIN motivo m ON m.id = p.motivo_id
        WHERE p.inicio < ? AND (p.fim IS NULL OR p.fim > ?)""", (fim.isoformat(), ini.isoformat()))

    por_eq = {e["id"]: {"equipamento": e["nome"], "min_parado": 0.0, "falhas": 0, "min_falhas": 0.0}
              for e in eqs}
    for p in paradas:
        e = por_eq.get(p["equipamento_id"])
        if e is None:
            continue
        p_ini = max(datetime.fromisoformat(p["inicio"]), ini)
        p_fim = min(datetime.fromisoformat(p["fim"]) if p["fim"] else fim, fim)
        minutos = max(0.0, (p_fim - p_ini).total_seconds() / 60)
        e["min_parado"] += minutos
        if p["programado"] == 0 and datetime.fromisoformat(p["inicio"]) >= ini:
            e["falhas"] += 1
            e["min_falhas"] += minutos

    linhas, tot_op, tot_falhas, tot_min_falhas = [], 0.0, 0, 0.0
    for e in por_eq.values():
        horas_op = max(0.0, horas_periodo - e["min_parado"] / 60)
        tot_op += horas_op
        tot_falhas += e["falhas"]
        tot_min_falhas += e["min_falhas"]
        linhas.append({
            "equipamento": e["equipamento"],
            "horas_operacao": round(horas_op, 2),
            "horas_paradas": round(e["min_parado"] / 60, 2),
            "falhas": e["falhas"],
            "mtbf_h": round(horas_op / e["falhas"], 2) if e["falhas"] else None,
            "mttr_min": round(e["min_falhas"] / e["falhas"], 2) if e["falhas"] else None,
        })
    return {
        "dias": dias,
        "desde": ini.isoformat(),
        "ate": fim.isoformat(),
        "geral": {
            "horas_operacao": round(tot_op, 2),
            "falhas": tot_falhas,
            "mtbf_h": round(tot_op / tot_falhas, 2) if tot_falhas else None,
            "mttr_min": round(tot_min_falhas / tot_falhas, 2) if tot_falhas else None,
        },
        "por_equipamento": linhas,
    }


@bp.get("/analytics/mtbf")
def mtbf():
    return jsonify(_mtbf(request.args))


# -------------------------------------------------------- Export Excel (v2.1)

AZUL = PatternFill("solid", fgColor="2B5FAB")
BRANCO_NEGRITO = Font(bold=True, color="FFFFFF")


def _local(iso):
    """Converte o horário salvo (UTC) para o horário de Brasília, sem fuso (Excel não aceita fuso)."""
    if not iso:
        return None
    return datetime.fromisoformat(iso).astimezone(FUSO).replace(tzinfo=None)


def _aba(wb, titulo, cabecalho, linhas, formatos=None, primeira=False):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = titulo
    ws.append(cabecalho)
    for c in ws[1]:
        c.fill, c.font = AZUL, BRANCO_NEGRITO
        c.alignment = Alignment(horizontal="center", vertical="center")
    for linha in linhas:
        ws.append(linha)
    for col, fmt in (formatos or {}).items():
        for c in ws[col][1:]:
            c.number_format = fmt
    for i, nome in enumerate(cabecalho, 1):
        letra = get_column_letter(i)
        maior = max([len(str(nome))] + [len(str(c.value)) for c in ws[letra][1:] if c.value is not None])
        ws.column_dimensions[letra].width = min(max(12, maior + 2), 50)
    ws.freeze_panes = "A2"
    if linhas:
        ws.auto_filter.ref = ws.dimensions
    return ws


@bp.get("/export.xlsx")
def exportar_excel():
    """Baixa a base em Excel: Paradas, MTBF por máquina, Pareto e Resumo."""
    paradas = _filtrar_paradas(request.args)
    m = _mtbf(request.args)
    data = "dd/mm/yyyy hh:mm:ss"

    wb = Workbook()
    _aba(wb, "Paradas",
         ["ID", "Equipamento", "Operador", "Início", "Fim", "Duração (min)", "Código motivo",
          "Motivo", "Categoria", "Tipo", "Observação", "Status"],
         [[p["id"], p["equipamento"], p["operador"], _local(p["inicio"]), _local(p["fim"]),
           p["duracao_min"], p["motivo_codigo"], p["motivo"], p["categoria"],
           None if p["programado"] is None else ("Programado" if p["programado"] else "Não programado"),
           p["observacao"], p["status"]] for p in paradas],
         {"D": data, "E": data, "F": "0.00"}, primeira=True)

    linhas = [[l["equipamento"], l["horas_operacao"], l["horas_paradas"], l["falhas"],
               l["mtbf_h"], l["mttr_min"]] for l in m["por_equipamento"]]
    g = m["geral"]
    linhas.append(["TOTAL", g["horas_operacao"], None, g["falhas"], g["mtbf_h"], g["mttr_min"]])
    ws = _aba(wb, "MTBF", ["Equipamento", "Horas em operação", "Horas paradas", "Falhas",
                           "MTBF (h)", "MTTR falhas (min)"], linhas,
              {"B": "0.00", "C": "0.00", "E": "0.00", "F": "0.00"})
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws.append([])
    ws.append([f"Período: últimos {m['dias']} dias ({_local(m['desde']):%d/%m/%Y %H:%M} a "
               f"{_local(m['ate']):%d/%m/%Y %H:%M})"])
    ws.append(["MTBF = horas em operação ÷ nº de falhas (paradas não programadas). "
               "Vazio = sem falhas no período."])

    _aba(wb, "Pareto", ["Código", "Motivo", "Tipo", "Ocorrências", "Minutos", "% acumulado"],
         [[x["codigo"], x["descricao"], "Programado" if x["programado"] else "Não programado",
           x["ocorrencias"], x["minutos"], x["pct_acumulado"] / 100] for x in _pareto()],
         {"E": "0.00", "F": "0.0%"})

    r = _resumo()
    _aba(wb, "Resumo", ["Indicador", "Valor"], [
        ["Paradas fechadas", r["total_paradas"]],
        ["Minutos parados", r["minutos_total"]],
        ["MTTR geral (min)", r["mttr_min"]],
        [f"MTBF últimos {m['dias']} dias (h)", g["mtbf_h"]],
        ["Minutos programados", r["min_programado"]],
        ["Minutos não programados", r["min_nao_programado"]],
        ["Paradas abertas / sem motivo", r["paradas_abertas"]],
        ["Gerado em", datetime.now(FUSO).replace(tzinfo=None, microsecond=0)],
        ["Versão do sistema", __version__],
    ])
    wb["Resumo"]["B9"].number_format = data

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"paradas_{datetime.now(FUSO):%Y-%m-%d_%H%M}.xlsx"
    return send_file(buf, as_attachment=True, download_name=nome,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
